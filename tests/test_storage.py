import sys
from datetime import datetime

import pytest
from openpyxl import load_workbook

from fire_inspector import storage
from fire_inspector.models import Record
from fire_inspector.storage import RecordStore

HEADERS = [
    "日期時間",
    "上傳者",
    "上傳者帳號ID",
    "衣服/褲子判定",
    "鞋子判定",
    "頭髮判定",
    "首飾/口罩/眼鏡判定",
    "綜合結果",
    "累計合格總次數",
    "照片存放位置",
]


def make_record(**overrides):
    fields = dict(
        taken_at=datetime(2026, 10, 27, 19, 30, 45),
        uploader="小明",
        user_id=123456789012345678,
        clothing="⚠️ 材質無法判斷",
        shoes="✅ 布鞋",
        hair="❌ 瀏海沒夾",
        other="首飾 ✅;口罩 ✅;眼鏡 ⚠️ 建議換隱形",
        overall="不合格",
        cumulative=7,
        photo_path="photos/20261027-193045_小明.jpg",
    )
    fields.update(overrides)
    return Record(**fields)


def locked(*args, **kwargs):
    raise PermissionError("Excel 正開著這個檔案")


# --- Excel 讀寫 ---


def test_first_add_creates_workbook_with_header_row(tmp_path):
    RecordStore(tmp_path).add(make_record())
    sheet = load_workbook(tmp_path / "records.xlsx").active
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.max_row == 2


def test_add_appends_instead_of_overwriting(tmp_path):
    store = RecordStore(tmp_path)
    store.add(make_record(uploader="甲"))
    store.add(make_record(uploader="乙"))
    store.add(make_record(uploader="丙"))
    assert [r.uploader for r in store.all_records()] == ["甲", "乙", "丙"]


def test_record_survives_a_round_trip_including_18_digit_user_id(tmp_path):
    original = make_record()
    RecordStore(tmp_path).add(original)
    assert RecordStore(tmp_path).all_records() == [original]


def test_all_records_is_empty_when_no_file_yet(tmp_path):
    assert RecordStore(tmp_path).all_records() == []


def test_manual_edit_in_excel_is_seen_on_next_read(tmp_path):
    store = RecordStore(tmp_path)
    store.add(make_record(overall="合格"))
    workbook = load_workbook(tmp_path / "records.xlsx")
    workbook.active["H2"] = "不合格"
    workbook.save(tmp_path / "records.xlsx")
    assert store.all_records()[0].overall == "不合格"


def test_formula_like_text_is_stored_as_plain_text(tmp_path):
    RecordStore(tmp_path).add(make_record(uploader='=HYPERLINK("http://x.test","x")'))
    cell = load_workbook(tmp_path / "records.xlsx").active["B2"]
    assert cell.data_type == "s"
    assert cell.value.startswith("'")


def test_no_temporary_file_is_left_behind(tmp_path):
    RecordStore(tmp_path).add(make_record())
    assert sorted(p.name for p in tmp_path.iterdir()) == ["records.xlsx"]


# --- Excel 被鎖住 ---


def test_locked_excel_does_not_lose_the_record(tmp_path, monkeypatch):
    store = RecordStore(tmp_path)
    store.add(make_record(uploader="甲"))
    monkeypatch.setattr(storage, "_append_excel", locked)

    store.add(make_record(uploader="乙"))  # 不可以丟出錯誤

    assert [r.uploader for r in store.all_records()] == ["甲", "乙"]
    assert (tmp_path / "pending.jsonl").read_text(encoding="utf-8").strip() != ""


def test_pending_records_are_written_once_excel_is_free(tmp_path, monkeypatch):
    store = RecordStore(tmp_path)
    store.add(make_record(uploader="甲"))
    with monkeypatch.context() as m:
        m.setattr(storage, "_append_excel", locked)
        store.add(make_record(uploader="乙"))

    store.flush_pending()

    assert [r.uploader for r in RecordStore(tmp_path).all_records()] == ["甲", "乙"]
    assert not (tmp_path / "pending.jsonl").exists()


def test_flush_while_still_locked_keeps_pending(tmp_path, monkeypatch):
    store = RecordStore(tmp_path)
    monkeypatch.setattr(storage, "_append_excel", locked)
    store.add(make_record(uploader="乙"))

    store.flush_pending()

    assert (tmp_path / "pending.jsonl").exists()
    assert [r.uploader for r in store.all_records()] == ["乙"]


def test_adding_after_unlock_flushes_older_pending_first_without_duplicates(
    tmp_path, monkeypatch
):
    store = RecordStore(tmp_path)
    with monkeypatch.context() as m:
        m.setattr(storage, "_append_excel", locked)
        store.add(make_record(uploader="乙"))
    store.add(make_record(uploader="丙"))

    assert [r.uploader for r in store.all_records()] == ["乙", "丙"]


def test_locked_read_falls_back_to_last_snapshot_plus_pending(tmp_path, monkeypatch):
    store = RecordStore(tmp_path)
    store.add(make_record(uploader="甲"))
    store.all_records()  # 成功讀過一次,留下快照
    monkeypatch.setattr(storage, "_read_excel", locked)
    monkeypatch.setattr(storage, "_append_excel", locked)
    store.add(make_record(uploader="乙"))

    assert [r.uploader for r in store.all_records()] == ["甲", "乙"]


def test_locked_read_without_any_snapshot_raises(tmp_path, monkeypatch):
    RecordStore(tmp_path).add(make_record())
    monkeypatch.setattr(storage, "_read_excel", locked)
    with pytest.raises(PermissionError):
        RecordStore(tmp_path).all_records()


# --- 照片存檔 ---


def test_save_photo_writes_into_photos_folder_and_returns_relative_path(tmp_path):
    relative = RecordStore(tmp_path).save_photo("20261027-193045_小明.jpg", b"jpeg-bytes")
    assert relative == "photos/20261027-193045_小明.jpg"
    assert (tmp_path / relative).read_bytes() == b"jpeg-bytes"


def test_save_photo_never_overwrites_same_second_collision(tmp_path):
    store = RecordStore(tmp_path)
    first = store.save_photo("a.jpg", b"1")
    second = store.save_photo("a.jpg", b"2")
    assert (first, second) == ("photos/a.jpg", "photos/a-2.jpg")
    assert (tmp_path / first).read_bytes() == b"1"


def test_save_photo_refuses_a_name_that_escapes_the_photos_folder(tmp_path):
    with pytest.raises(ValueError):
        RecordStore(tmp_path).save_photo("../records.xlsx", b"evil")
    assert not (tmp_path / "records.xlsx").exists()


# --- 第一次啟動、真實的檔案鎖 ---


def test_add_creates_the_data_folder_when_it_does_not_exist_yet(tmp_path):
    data_dir = tmp_path / "not-created-yet"
    RecordStore(data_dir).add(make_record())
    assert (data_dir / "records.xlsx").exists()


@pytest.mark.skipif(sys.platform != "win32", reason="只有 Windows 會鎖住開著的檔案")
def test_real_windows_file_lock_sends_the_record_to_pending(tmp_path):
    store = RecordStore(tmp_path)
    store.add(make_record(uploader="甲"))

    with open(tmp_path / "records.xlsx", "rb"):  # 模擬 Excel 開著這個檔案
        store.add(make_record(uploader="乙"))

    assert (tmp_path / "pending.jsonl").exists()
    store.flush_pending()
    assert [r.uploader for r in RecordStore(tmp_path).all_records()] == ["甲", "乙"]
