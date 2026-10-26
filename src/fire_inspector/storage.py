"""Excel 紀錄、Excel 被鎖住時的暫存區、照片存檔。"""

import json
import os
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from fire_inspector.models import Record
from fire_inspector.sanitize import ensure_inside, safe_cell_text, unique_path

HEADERS = [
    "日期時間",
    "上傳者",
    "上傳者帳號ID",
    "衣服/褲子判定",
    "鞋子判定",
    "頭髮判定",
    "首飾/口罩/眼鏡判定",
    "綜合結果",
    "累計合格總次數",
    "照片存放位置",
]
_TIME_FORMAT = "%Y-%m-%d %H:%M:%S"


def _to_row(record: Record) -> list:
    return [
        record.taken_at.strftime(_TIME_FORMAT),
        safe_cell_text(record.uploader),
        # Discord ID 有 18 位數,Excel 數字只保證 15 位精確,所以用文字存
        str(record.user_id),
        safe_cell_text(record.clothing),
        safe_cell_text(record.shoes),
        safe_cell_text(record.hair),
        safe_cell_text(record.other),
        safe_cell_text(record.overall),
        record.cumulative,
        safe_cell_text(record.photo_path),
    ]


def _from_row(row: tuple) -> Record:
    taken_at = row[0]
    if not isinstance(taken_at, datetime):
        taken_at = datetime.strptime(str(taken_at), _TIME_FORMAT)

    def text(index: int) -> str:
        return "" if row[index] is None else str(row[index])

    return Record(
        taken_at=taken_at,
        uploader=text(1),
        user_id=int(str(row[2]).strip()),
        clothing=text(3),
        shoes=text(4),
        hair=text(5),
        other=text(6),
        overall=text(7),
        cumulative=int(row[8] or 0),
        photo_path=text(9),
    )


def _read_excel(path: Path) -> list[Record]:
    if not path.exists():
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(min_row=2, values_only=True)
        return [_from_row(row) for row in rows if row and row[0] is not None]
    finally:
        workbook.close()


def _append_excel(path: Path, records: list[Record]) -> None:
    """在最後面加上新的列。先寫暫存檔再換掉正本,避免寫到一半當機而毀檔。"""
    if path.exists():
        workbook = load_workbook(path)
    else:
        workbook = Workbook()
        workbook.active.append(HEADERS)
    for record in records:
        workbook.active.append(_to_row(record))

    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_name(path.name + ".tmp")
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        temp_path.unlink(missing_ok=True)


class RecordStore:
    def __init__(self, data_dir: Path):
        self.data_dir = Path(data_dir)
        self.excel_path = self.data_dir / "records.xlsx"
        self.pending_path = self.data_dir / "pending.jsonl"
        self.photos_dir = self.data_dir / "photos"
        # Excel 內容的快照:最近一次成功讀取,加上之後成功寫入的列
        self._snapshot: list[Record] | None = None

    def add(self, record: Record) -> None:
        """寫進 Excel。Excel 被鎖住時改存暫存區,不丟出錯誤。"""
        pending = self._read_pending()
        try:
            _append_excel(self.excel_path, [*pending, record])
        except PermissionError:
            self._write_pending([*pending, record])
            return
        self._after_written([*pending, record])

    def flush_pending(self) -> None:
        pending = self._read_pending()
        if not pending:
            return
        try:
            _append_excel(self.excel_path, pending)
        except PermissionError:
            return
        self._after_written(pending)

    def all_records(self) -> list[Record]:
        """Excel 的內容加上暫存區。Excel 被鎖住就用快照;連快照都沒有時丟出 PermissionError。"""
        pending = self._read_pending()
        try:
            self._snapshot = _read_excel(self.excel_path)
        except PermissionError:
            if self._snapshot is None:
                raise
        return [*self._snapshot, *pending]

    def save_photo(self, filename: str, data: bytes) -> str:
        """存照片,回傳相對於 DATA_DIR 的路徑。同名不覆蓋,路徑跑出資料夾就拒絕。"""
        self.photos_dir.mkdir(parents=True, exist_ok=True)
        path = unique_path(self.photos_dir, filename)
        ensure_inside(self.photos_dir, path)
        path.write_bytes(data)
        return path.relative_to(self.data_dir).as_posix()

    def _after_written(self, written: list[Record]) -> None:
        self.pending_path.unlink(missing_ok=True)
        if self._snapshot is not None:
            self._snapshot = [*self._snapshot, *written]

    def _read_pending(self) -> list[Record]:
        if not self.pending_path.exists():
            return []
        records = []
        for line in self.pending_path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                data = json.loads(line)
                data["taken_at"] = datetime.fromisoformat(data["taken_at"])
                records.append(Record(**data))
        return records

    def _write_pending(self, records: list[Record]) -> None:
        self.data_dir.mkdir(parents=True, exist_ok=True)
        lines = []
        for record in records:
            data = asdict(record)
            data["taken_at"] = record.taken_at.isoformat()
            lines.append(json.dumps(data, ensure_ascii=False))
        self.pending_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
